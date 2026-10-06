from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('output')
out.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('Halaman '); r.font.size = Pt(9)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)
def p(text, bold=False):
    x=doc.add_paragraph(); x.add_run(text).bold=bold; return x
def h(text): doc.add_heading(text, 2)
def item(label, text):
    x=doc.add_paragraph(); x.add_run(label+' ').bold=True; x.add_run(text)
def page(title):
    doc.add_page_break(); doc.add_heading(title, 1)

doc.add_paragraph('Kegawatan Pankreatitis Akut', 'Title')
p('Pengenalan tanda bahaya dan tindakan awal keperawatan pada pasien dewasa')
p('Pembaruan sumber 18 September 2026')
doc.add_heading('Prioritas keselamatan pasien',1)
p('Kegawatan pankreatitis akut terutama berupa gangguan sirkulasi, pernapasan, ginjal, infeksi berat, dan komplikasi intraabdomen. Kondisi dapat memburuk cepat. Lakukan penilaian ABCDE dan eskalasi segera saat ditemukan gangguan organ, tanpa menunggu hasil pemeriksaan lengkap. [1–4]')
p('Gagal organ yang berlangsung lebih dari 48 jam menentukan pankreatitis akut berat. Batas ini digunakan untuk klasifikasi, bukan untuk menunda pertolongan atau rujukan ICU. Nekrosis atau koleksi cairan saja tidak selalu merupakan kegawatan. [4]')
h('Tanda yang memerlukan respons segera')
for s in [
'Sesak baru atau memburuk, penurunan saturasi, kebutuhan oksigen meningkat, atau napas tidak adekuat.',
'Tekanan darah turun, ekstremitas dingin, pengisian kapiler memanjang, nadi cepat, pingsan, atau penurunan kesadaran.',
'Urine turun nyata atau tidak keluar, terutama disertai hipotensi atau sesak.',
'Abdomen semakin tegang, nyeri berubah mendadak, tanda peritonitis, atau perdarahan.',
'Demam atau hipotermia disertai perburukan klinis, kebingungan, atau gangguan perfusi.',
'Kejang, aritmia, tetani, atau perubahan kesadaran yang mengarah pada gangguan metabolik.'
]: doc.add_paragraph(s,'List Bullet')
h('Panduan penggunaan')
p('Dokumen ini ditujukan untuk pembelajaran dan praktik keperawatan di fasilitas kesehatan. Tindakan mandiri meliputi pengkajian, pemantauan, perlindungan keselamatan, dokumentasi, dan eskalasi. Cairan, obat, oksigen, pemeriksaan, serta tindakan invasif dilaksanakan menurut kewenangan, protokol, dan instruksi tim medis. Frekuensi observasi mengikuti kondisi pasien dan sistem peringatan dini rumah sakit.')
p('Apabila tanda bahaya muncul di rumah, pasien memerlukan pertolongan IGD segera.', True)

page('Kegawatan sirkulasi pernapasan dan ginjal')
h('1 Syok dan hipoperfusi')
item('Gambaran klinis', 'Kehilangan cairan melalui muntah dan perpindahan cairan ke jaringan dapat mengurangi volume sirkulasi. Syok juga dapat terjadi karena sepsis atau perdarahan. Waspadai hipotensi, kulit dingin atau mottling, nadi lemah, kebingungan, dan oliguria.')
item('Tindakan awal perawat', 'Aktifkan bantuan, kaji ABCDE, pantau tekanan darah dan perfusi berulang, pastikan akses intravena, serta catat cairan masuk dan keluar. Periksa adanya tanda kelebihan cairan selama terapi.')
item('Kolaborasi', 'Pemeriksaan laktat dan fungsi organ, kristaloid yang dititrasi terhadap respons, serta vasopresor bila diperlukan. MAP sekitar 65 mmHg sering menjadi target awal syok septik, dengan penyesuaian individual. [5]')
p('Resusitasi cairan tidak boleh menggunakan volume tetap tanpa evaluasi. Strategi cairan agresif rutin pada pankreatitis meningkatkan risiko kelebihan cairan; pasien syok tetap memerlukan resusitasi segera yang disesuaikan. [2,6]')
h('2 Gagal napas akut termasuk ARDS')
item('Gambaran klinis', 'Sesak, takipnea, peningkatan usaha napas, saturasi turun, atau penurunan kesadaran. Penyebab dapat mencakup acute respiratory distress syndrome (ARDS), edema paru, efusi pleura, aspirasi, atau depresi napas akibat obat.')
item('Tindakan awal perawat', 'Pastikan jalan napas, posisikan untuk membantu ventilasi bila hemodinamik memungkinkan, pantau SpO₂ dan kesadaran, serta berikan oksigen sesuai protokol. Siapkan pengisapan bila ada muntahan dan bantuan ventilasi bila diperlukan.')
item('Kolaborasi', 'Evaluasi gas darah, penyebab hipoksemia, dan kebutuhan dukungan napas atau ICU. Jangan menganggap setiap sesak disebabkan nyeri atau kecemasan. [3]')
h('3 Cedera ginjal akut')
item('Gambaran klinis', 'Urine berkurang, kreatinin meningkat, edema, atau gangguan kalium dan asam basa. Produksi urine <0,5 mL/kg/jam selama sedikitnya 6 jam memenuhi salah satu kriteria AKI, tetapi penurunan mendadak harus dilaporkan lebih awal. [7]')
item('Tindakan awal perawat', 'Ukur keluaran urine secara akurat, tinjau tren cairan, dan periksa selang kateter bila terpasang. Laporkan perubahan; jangan otomatis menambah cairan hanya berdasarkan oliguria.')
item('Kolaborasi', 'Evaluasi perfusi, obstruksi, obat nefrotoksik, elektrolit, dan kebutuhan terapi pengganti ginjal. Anuria, hiperkalemia dengan perubahan EKG, edema paru, atau asidosis berat memerlukan evaluasi segera.')

page('Infeksi berat dan kegawatan intraabdomen')
h('4 Sepsis dan syok septik')
item('Gambaran klinis', 'Infeksi dapat berasal dari nekrosis terinfeksi, saluran empedu, paru, urine, atau alat invasif. Curigai bila demam atau hipotermia disertai penurunan kesadaran, hipotensi, oliguria, atau peningkatan kebutuhan oksigen. Demam dan leukositosis saja dapat mencerminkan inflamasi steril.')
item('Tindakan awal perawat', 'Aktifkan alur sepsis, catat waktu pengenalan, lakukan pemantauan ketat, dan fasilitasi kultur serta laktat. Kultur diambil sebelum antibiotik bila memungkinkan tanpa menunda terapi secara bermakna.')
item('Kolaborasi', 'Pada dugaan syok septik, antibiotik diberikan segera, idealnya dalam satu jam pengenalan; lakukan evaluasi pengendalian sumber infeksi. Antibiotik profilaksis tidak diberikan rutin untuk nekrosis steril. [2,5]')
h('5 Perdarahan berat')
item('Gambaran klinis', 'Perdarahan dapat terjadi akibat erosi pembuluh atau ruptur pseudoaneurisma. Waspadai nyeri mendadak, pucat, hipotensi, penurunan hemoglobin, hematemesis, melena, atau darah dari drain. Perdarahan internal dapat terjadi tanpa darah yang tampak.')
item('Tindakan awal perawat', 'Aktifkan respons kegawatdaruratan, nilai perfusi, pertahankan akses intravena, ukur kehilangan darah yang tampak, dan siapkan pemeriksaan serta transfusi sesuai instruksi. Jangan menunggu penurunan Hb untuk melaporkan syok.')
item('Kolaborasi', 'Evaluasi radiologi intervensi untuk embolisasi dan/atau pembedahan. Perdarahan aktif dengan instabilitas tidak mengikuti prinsip penundaan intervensi pada nekrosis yang stabil. [1,3]')
h('6 Sindrom kompartemen abdomen')
item('Gambaran klinis', 'Abdomen tegang membesar, oliguria, ventilasi memburuk, dan hipotensi. Definisi: tekanan intraabdomen menetap >20 mmHg disertai disfungsi organ baru. Distensi atau lingkar perut saja tidak menegakkan diagnosis. [8]')
item('Tindakan awal perawat', 'Laporkan perubahan abdomen bersama fungsi napas dan ginjal; bantu pengukuran tekanan intraabdomen sesuai protokol. Pantau respons dan keseimbangan cairan.')
item('Kolaborasi', 'Dekompresi gastrointestinal, drainase, optimasi cairan dan analgesia, serta dekompresi bedah bila terapi awal gagal. [3]')
h('7 Iskemia atau perforasi usus')
item('Gambaran klinis', 'Nyeri baru yang sangat berat, peritonitis, distensi progresif, perdarahan saluran cerna, atau perburukan syok.')
item('Tindakan awal dan kolaborasi', 'Minta evaluasi bedah segera, lakukan pemantauan ABCDE, dan hentikan pemberian enteral sementara bila dicurigai iskemia/perforasi sambil menghubungi tim medis. Siapkan pemeriksaan serta tindakan sesuai instruksi. [1,3]')

page('Kegawatan lain dan alur respons perawat')
h('8 Kolangitis akut pada pankreatitis bilier')
p('Obstruksi bilier yang terinfeksi dapat menyebabkan sepsis. Waspadai demam atau menggigil, ikterus, nyeri abdomen, terutama bila disertai hipotensi atau kebingungan. Kombinasi gejala klasik tidak harus lengkap. Laporkan segera, lakukan pemantauan sepsis, dan siapkan antibiotik serta drainase bilier sesuai instruksi. Pankreatitis dengan kolangitis memerlukan ERCP segera; ACG mendukung pelaksanaan dalam 24 jam. ERCP dini tidak rutin untuk pankreatitis bilier tanpa kolangitis. [2]')
h('9 Gangguan elektrolit glukosa dan asam basa yang berat')
p('Hipokalsemia simptomatik dapat menimbulkan tetani atau kejang; gangguan kalium dapat menyebabkan aritmia. Hipoglikemia, krisis hiperglikemik, dan asidosis berat juga dapat mengancam jiwa. Kaji kesadaran, glukosa, gejala neuromuskular, dan EKG sesuai indikasi. Kolaborasikan pemeriksaan serta koreksi terukur menurut protokol; jangan memberikan koreksi elektrolit cepat tanpa pemantauan yang diperlukan.')
h('10 Gagal multiorgan')
p('Gangguan beberapa organ sekaligus, misalnya syok bersama gagal napas dan cedera ginjal, menandakan kondisi sangat berat. Perawat memprioritaskan stabilisasi ABCDE, koordinasi ICU, serta pemantauan respons terapi. Jangan menunggu klasifikasi gagal organ persisten untuk meminta bantuan. [3,4]')
h('Alur tindakan saat ditemukan perburukan')
for s in [
'Panggil bantuan atau tim respons cepat. Bila henti jantung, aktifkan kode emergensi dan lakukan resusitasi sesuai protokol.',
'Lakukan ABCDE secara berurutan; tangani ancaman nyawa yang ditemukan dan ulangi penilaian setelah tindakan.',
'Pasang pemantauan yang diperlukan, pastikan akses intravena, nilai kesadaran, dan ukur glukosa bila relevan.',
'Siapkan data tren tanda vital, kebutuhan oksigen, urine, keseimbangan cairan, obat, serta hasil laboratorium.',
'Laksanakan instruksi kolaboratif dan nilai responsnya. Tingkatkan level perawatan bila kebutuhan pemantauan atau dukungan organ melebihi kemampuan unit.',
'Dokumentasikan waktu perburukan, siapa yang dihubungi, tindakan, dan respons pasien.'
]: doc.add_paragraph(s,'List Number')
h('Komunikasi SBAR')
p('S: perubahan yang paling mengancam. B: onset pankreatitis, penyebab, komorbiditas, dan terapi. A: ABCDE serta tren perfusi, napas, urine, dan laboratorium. R: minta penilaian segera, instruksi stabilisasi, serta keputusan pemindahan ke unit yang sesuai.')
p('Nyeri berat dan muntah memerlukan penanganan cepat, tetapi juga harus dinilai sebagai kemungkinan tanda komplikasi. Jangan hanya meningkatkan analgesik tanpa mengevaluasi perubahan klinis.')

page('Daftar pustaka')
p('Sumber pedoman terbaru dilengkapi konsensus dasar yang masih relevan. Uraian tindakan awal merupakan sintesis keperawatan dan perlu diintegrasikan dengan protokol fasilitas.')
refs=[
('1','Cárdenas-Jaén K, dkk. AEG-AESPANC-OPGE-SIED-SPG Ibero-Latin American Guidelines on Acute Pancreatitis. United European Gastroenterology Journal. 2026.','https://doi.org/10.1002/ueg2.70190'),
('2','Tenner S, dkk. American College of Gastroenterology Guidelines: Management of Acute Pancreatitis. American Journal of Gastroenterology. 2024;119:419–437.','https://doi.org/10.14309/ajg.0000000000002645'),
('3','Leppäniemi A, dkk. 2019 WSES guidelines for the management of severe acute pancreatitis. World Journal of Emergency Surgery. 2019;14:27.','https://doi.org/10.1186/s13017-019-0247-0'),
('4','Banks PA, dkk. Classification of acute pancreatitis—2012: revision of the Atlanta classification and definitions by international consensus. Gut. 2013;62:102–111.','https://doi.org/10.1136/gutjnl-2012-302779'),
('5','Society of Critical Care Medicine. Surviving Sepsis Campaign International Guidelines for Management of Sepsis and Septic Shock 2026.','https://www.sccm.org/survivingsepsiscampaign/guidelines-and-resources/surviving-sepsis-campaign-adult-guidelines'),
('6','de-Madaria E, dkk. Aggressive or Moderate Fluid Resuscitation in Acute Pancreatitis. New England Journal of Medicine. 2022;387:989–1000.','https://doi.org/10.1056/NEJMoa2202884'),
('7','KDIGO. Clinical Practice Guideline for Acute Kidney Injury. Kidney International Supplements. 2012;2:1–138.','https://kdigo.org/wp-content/uploads/2016/10/KDIGO-2012-AKI-Guideline-English.pdf'),
('8','World Society of the Abdominal Compartment Syndrome. WSACS Consensus Guidelines Summary.','https://www.wsacs.org/education/436/wsacs-consensus-guidelines-summary/'),
]
for n,t,url in refs:
    x=p(f'[{n}] {t}')
    x.paragraph_format.keep_with_next=True
    x=p(url)
    for r in x.runs: r.font.size=Pt(9)
p('Diakses 18 September 2026. Kriteria AKI menggunakan pedoman KDIGO final; draf pembaruan tidak diperlakukan sebagai standar final.')
doc.core_properties.title='Kegawatan Pankreatitis Akut'
doc.core_properties.subject='Tanda bahaya dan tindakan awal keperawatan pada pasien dewasa'
doc.core_properties.author=''
doc.save(out/'Kegawatan_Pankreatitis_Akut.docx')
print(out.resolve()/'Kegawatan_Pankreatitis_Akut.docx')
