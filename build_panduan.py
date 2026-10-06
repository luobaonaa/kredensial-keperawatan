from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.8)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=d.styles[n]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(24)
for n,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    d.styles[n].font.size=Pt(size)
    d.styles[n].paragraph_format.space_before=Pt(10)
    d.styles[n].paragraph_format.space_after=Pt(5)
f=s.footer.paragraphs[0]; f.alignment=2
f.add_run('Halaman ').font.size=Pt(9)
el=OxmlElement('w:fldSimple'); el.set(qn('w:instr'),'PAGE'); f._p.append(el)
def p(t): return d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def label(a,b):
    x=d.add_paragraph(); x.add_run(a+' ').bold=True; x.add_run(b)
def bullets(items):
    for t in items: d.add_paragraph(t,'List Bullet')
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers,rows,widths):
    t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,v in enumerate(headers): t.rows[0].cells[i].text=v
    for vals in rows:
        for c,v in zip(t.add_row().cells,vals): c.text=v
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
    pr.append(borders)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=1
            cp=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge); m.set(qn('w:w'),'90'); m.set(qn('w:type'),'dxa'); margins.append(m)
            cp.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs: r.font.size=Pt(10)
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF2'); cp.append(sh)
                for r in c.paragraphs[0].runs: r.bold=True
        nr=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(nr)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p('')

d.add_paragraph('Panduan Asuhan Keperawatan', 'Title')
d.add_paragraph('Pasien Dewasa dengan Pankreatitis Akut','Subtitle')
p('Lingkup IGD, ruang rawat inap, dan perawatan intensif\nDasar pembaruan sumber 18 September 2026')
h('1 Tujuan dan ketentuan penggunaan')
p('Panduan ini membantu perawat mempertahankan fungsi organ, mengendalikan gejala, memenuhi kebutuhan nutrisi, mencegah komplikasi, serta mempersiapkan pasien dan keluarga untuk perawatan setelah pulang. Prioritas utama adalah mengenali perburukan dan melakukan eskalasi segera.')
p('Penerapan disesuaikan dengan derajat penyakit, kondisi awal, komorbiditas, kewenangan klinis, instruksi dokter penanggung jawab pelayanan (DPJP), dan prosedur rumah sakit. Anak, ibu hamil, serta pasien dengan gangguan jantung atau ginjal memerlukan penyesuaian khusus.')
p('Diagnosis keperawatan ditetapkan berdasarkan hasil pengkajian; tidak semua diagnosis di bawah harus digunakan. Nama, kode, indikator diagnostik, dan skala luaran dicocokkan dengan SDKI, SLKI, dan SIKI resmi yang berlaku di institusi. Rincian tindakan berikut adalah sintesis operasional keperawatan, bukan kutipan lengkap standar PPNI.')
h('2 Konsep klinis')
p('Pankreatitis akut adalah peradangan akut pankreas yang dapat menyebabkan nyeri, gangguan pencernaan, perpindahan cairan ke jaringan, dan gangguan organ. Diagnosis medis memerlukan sedikitnya dua dari tiga temuan: nyeri abdomen khas, lipase atau amilase lebih dari tiga kali batas atas normal, atau pencitraan yang mendukung. Penyebab meliputi batu empedu, alkohol, hipertrigliseridemia, obat, hiperkalsemia, dan kondisi lain. [1,9]')
table(['Derajat','Kriteria utama'],[
('Ringan','Tanpa gagal organ maupun komplikasi lokal atau sistemik.'),
('Sedang sampai berat','Gagal organ sementara yang pulih dalam 48 jam dan/atau komplikasi tanpa gagal organ persisten.'),
('Berat','Gagal satu atau lebih organ yang menetap lebih dari 48 jam.')],[1.5,5.4])
p('Klasifikasi mengikuti Revised Atlanta. Jangan menunggu 48 jam untuk menangani atau merujuk pasien yang memburuk; batas tersebut digunakan untuk klasifikasi. [8]')

page('3 Pengkajian keperawatan')
h('Pengkajian awal ABCDE')
table(['Komponen','Data yang dinilai'],[
('A Airway','Kepatenan jalan napas, muntahan, kemampuan melindungi jalan napas, risiko aspirasi.'),
('B Breathing','Frekuensi dan pola napas, SpO₂, otot bantu, suara napas, kebutuhan oksigen.'),
('C Circulation','Tekanan darah, nadi, pengisian kapiler, suhu ekstremitas, urine, dehidrasi dan kelebihan cairan.'),
('D Disability','Kesadaran, kebingungan, glukosa sesuai indikasi, efek sedasi obat.'),
('E Exposure','Suhu, nyeri dan distensi abdomen, ikterus, edema, kulit, tanda perdarahan.')],[1.5,5.4])
h('Anamnesis terarah')
bullets([
'Kaji onset, lokasi, penjalaran, intensitas dan pencetus nyeri; gunakan skala yang sesuai kemampuan komunikasi pasien.',
'Catat mual, frekuensi dan perkiraan volume muntah, toleransi makan/minum, urine, buang air besar, dan flatus.',
'Tanyakan riwayat pankreatitis, batu empedu, ERCP, diabetes, penyakit jantung, ginjal, atau hati.',
'Identifikasi obat resep, obat bebas, suplemen, antikoagulan, dan alergi.',
'Kaji alkohol secara tidak menghakimi, termasuk konsumsi terakhir dan riwayat putus alkohol. Jangan otomatis menganggap alkohol sebagai penyebab.',
'Kaji berat badan sebelum sakit, penurunan berat badan, asupan terakhir, kemampuan aktivitas, dukungan keluarga, kecemasan, kebutuhan spiritual, dan hambatan komunikasi.'
])
h('Pemeriksaan fisik dan penunjang')
p('Nilai dampak nyeri terhadap tidur, napas dalam, gerak, dan makan. Kaji hidrasi, edema, abdomen, integritas kulit, risiko jatuh, luka tekan, dan kebutuhan bantuan aktivitas harian.')
p('Pantau hasil serta tren darah lengkap, hematokrit, ureum/BUN, kreatinin, elektrolit, kalsium, glukosa, fungsi hati, bilirubin, dan trigliserida sesuai program. Lipase/amilase membantu diagnosis. Gas darah dan laktat dipertimbangkan pada gangguan napas atau perfusi; kultur bila dicurigai infeksi.')
p('USG membantu evaluasi penyebab bilier. CT tidak rutin dilakukan saat masuk hanya untuk menilai keparahan; dipertimbangkan bila diagnosis belum jelas atau pasien tidak membaik dalam 48–72 jam. [4]')

page('4 Prioritas dan pemantauan')
p('Dahulukan gangguan jalan napas, pernapasan, perfusi, dan kesadaran. Setelah ancaman nyawa ditangani, prioritaskan nyeri, mual, nutrisi, mobilitas, serta kebutuhan psikososial. Gunakan sistem peringatan dini rumah sakit tanpa menggantikan penilaian klinis.')
p('Frekuensi berikut adalah contoh operasional. Tingkatkan observasi saat kondisi berubah, sesudah tindakan, atau sesuai skor peringatan dini.')
table(['Parameter','Pasien stabil','Pasien tidak stabil'],[
('Tanda vital dan kesadaran','Misalnya setiap 4 jam; lebih sering bila berisiko.','Pemantauan kontinu bila tersedia dan penilaian langsung berulang.'),
('Nyeri','Setiap observasi rutin dan saat keluhan berubah.','Sebelum dan sesudah terapi sesuai onset obat.'),
('Napas dan SpO₂','Bersama tanda vital.','Kontinu bila hipoksemia, dukungan meningkat, atau risiko depresi napas.'),
('Cairan dan urine','Setiap sif serta kumulatif 24 jam.','Urine per jam bila diperlukan secara klinis.'),
('Toleransi nutrisi','Setiap pemberian makan dan setiap sif.','Sesuai rute, laju, serta kondisi abdomen.'),
('Berat badan dan edema','Harian bila relevan dan memungkinkan.','Disesuaikan dengan keamanan pasien.'),
('Laboratorium','Menurut rencana DPJP.','Sesuai gangguan organ dan respons terapi.')],[1.35,2.5,3.05])
p('Kateter urine digunakan bila manfaat pengukuran akurat melebihi risikonya. Evaluasi kebutuhan setiap hari.')
h('Prinsip penetapan diagnosis dan luaran')
bullets([
'Diagnosis aktual memerlukan data yang memenuhi kriteria. Diagnosis risiko digunakan bila baru terdapat faktor risiko.',
'Tetapkan nilai awal, target terukur, dan waktu evaluasi. Target klinis disesuaikan kondisi serta preferensi pasien.',
'Pilih luaran SLKI dan intervensi SIKI yang sesuai setelah memeriksa buku resmi; dokumentasikan indikator yang benar-benar dinilai.',
'Intervensi mandiri meliputi observasi, kenyamanan, edukasi, pencegahan cedera, dan eskalasi. Pemberian obat, cairan, oksigen, serta tindakan invasif mengikuti kewenangan dan program kolaboratif.'
])

page('5 Rencana asuhan untuk cairan dan nyeri')
h('A Hipovolemia')
label('Data pendukung', 'Muntah, asupan menurun, mukosa kering, takikardia, hipotensi, perfusi perifer menurun, atau oliguria. Gunakan diagnosis risiko bila masalah belum terjadi.')
label('Luaran', 'Status cairan dan perfusi membaik pada evaluasi awal beberapa jam dan dipertahankan. Kesadaran serta perfusi membaik; tekanan darah mencapai target individual; keluaran urine umumnya ≥0,5 mL/kg/jam dengan interpretasi sesuai fungsi ginjal; tidak muncul kelebihan cairan.')
bullets([
'Catat semua cairan masuk, muntah, urine, dan keluaran drain.',
'Pantau nadi, tekanan darah, pengisian kapiler, edema, suara napas, serta kebutuhan oksigen.',
'Pastikan akses intravena berfungsi dan pemberian cairan sesuai instruksi.',
'Kolaborasikan kristaloid, biasanya Ringer laktat bila sesuai kondisi. Nilai kembali respons setelah perubahan terapi.',
'Laporkan sesak baru, ronki, edema progresif, atau kebutuhan oksigen meningkat. Jangan otomatis menambah cairan hanya karena oliguria.'
])
label('Rasional', 'Cairan dititrasi terhadap perfusi dan risiko kelebihan cairan. IAP 2025 mendukung strategi moderat; laju awal sekitar 1,5 mL/kg/jam merupakan pertimbangan medis, bukan resep seragam, terutama pada syok dan komorbiditas. [3]')
p('WATERFALL menunjukkan lebih banyak kelebihan cairan pada strategi agresif tanpa perbaikan luaran utama. Temuan ini tidak menggantikan resusitasi individual pada syok. [6]')
h('B Nyeri akut')
label('Data pendukung', 'Keluhan nyeri, meringis, sikap protektif, gelisah, atau aktivitas terbatas karena nyeri.')
label('Luaran', 'Nyeri menurun setelah intervensi, misalnya NRS ≤3/10 atau tingkat nyaman yang disepakati, disertai perbaikan istirahat dan gerak.')
bullets([
'Kaji karakter nyeri dan dampaknya terhadap fungsi; nilai ulang sesudah terapi.',
'Bantu posisi nyaman, misalnya semi-Fowler atau miring dengan lutut sedikit fleksi bila ditoleransi.',
'Ajarkan napas perlahan dan relaksasi; kurangi gangguan lingkungan.',
'Berikan analgesik sesuai resep. Pantau kesadaran, napas, tekanan darah, mual, serta fungsi usus setelah opioid.',
'Laporkan nyeri yang berubah mendadak atau memburuk meskipun sudah diterapi.'
])
label('Rasional', 'Analgesik mengikuti intensitas nyeri dan kondisi pasien. Opioid dapat diperlukan untuk nyeri berat. Bukti perbandingan obat masih terbatas; tidak ada regimen tunggal bagi semua pasien. [2]')

page('6 Rencana asuhan untuk mual dan nutrisi')
h('C Nausea')
label('Data pendukung', 'Keluhan mual, ingin muntah, salivasi meningkat, atau penolakan makan karena mual.')
label('Luaran', 'Mual berkurang, muntah berhenti atau menurun, dan asupan yang direncanakan dapat ditoleransi.')
bullets([
'Catat pencetus, frekuensi, volume, dan karakter muntah.',
'Berikan kebersihan mulut dan lingkungan bebas bau menyengat.',
'Atur posisi aman saat muntah; siapkan pengisapan bila diperlukan.',
'Kolaborasikan antiemetik dan koreksi cairan/elektrolit.',
'Evaluasi hubungan keluhan dengan obat dan nutrisi. Muntah menetap atau distensi memburuk memerlukan evaluasi medis.'
])
label('Rasional', 'Pengendalian mual mengurangi kehilangan cairan, meningkatkan kenyamanan, dan membantu pemulihan asupan.')
h('D Defisit nutrisi atau risiko defisit nutrisi')
label('Data pendukung', 'Hasil skrining, penurunan berat badan, dan asupan tidak adekuat yang memenuhi kriteria. Puasa singkat saja tidak otomatis menetapkan diagnosis aktual.')
label('Luaran', 'Tersedia rencana nutrisi dalam 24 jam, toleransi meningkat, dan kebutuhan dipenuhi bertahap sesuai penilaian ahli gizi.')
bullets([
'Lakukan skrining gizi dan catat persentase asupan aktual.',
'Kolaborasikan makan oral dini saat aman dan ditoleransi. Pada penyakit ringan, makanan lunak/padat rendah lemak dapat diberikan sejak awal; tidak harus melalui cair jernih.',
'Jangan mensyaratkan lipase normal sebelum memulai makan.',
'Bila oral tidak mencukupi, kolaborasikan enteral; pada pasien yang membutuhkan dukungan, umumnya dimulai dalam 24–72 jam.',
'Pantau muntah, distensi, diare, serta risiko aspirasi. Verifikasi posisi selang sesuai protokol dan tinggikan kepala tempat tidur 30–45° bila tidak kontraindikasi.',
'Pantau risiko refeeding syndrome pada malnutrisi atau asupan sangat rendah berkepanjangan; kolaborasikan pemantauan fosfat, kalium, magnesium, dan kebutuhan tiamin.',
'Parenteral dipertimbangkan jika enteral tidak memungkinkan, kontraindikasi, atau tidak mencukupi kebutuhan.'
])
label('Rasional', 'Saluran cerna diutamakan bila dapat digunakan. Rute dan peningkatan asupan mempertimbangkan stabilitas sirkulasi serta fungsi gastrointestinal. [5]')

page('7 Rencana asuhan untuk fungsi organ dan keselamatan')
h('E Gangguan pertukaran gas atau pola napas tidak efektif')
label('Data dan luaran', 'Pilih diagnosis berdasarkan hipoksemia/perubahan gas darah atau gangguan pola dan usaha napas. Targetnya adalah oksigenasi sesuai sasaran individual, usaha napas berkurang, dan kesadaran terjaga.')
label('Intervensi', 'Pantau SpO₂, frekuensi napas, suara napas, dan kebutuhan oksigen; posisikan untuk membantu ventilasi; kolaborasikan oksigen dengan target tertulis. Optimalkan kenyamanan untuk napas dalam dan mobilisasi bila stabil. Laporkan penurunan saturasi atau peningkatan dukungan; siapkan bantuan ventilasi sesuai instruksi.')
label('Rasional', 'Gangguan napas dapat menunjukkan komplikasi paru, kelebihan cairan, atau kegagalan organ. Disfungsi organ yang menetap memerlukan pemantauan intensif dan evaluasi ICU. [7]')
h('F Risiko ketidakseimbangan elektrolit dan ketidakstabilan glukosa')
label('Data dan luaran', 'Faktor terkait meliputi muntah, perubahan asupan, terapi cairan, fungsi ginjal, dan stres metabolik. Targetnya tidak ada manifestasi gangguan yang membahayakan dan hasil pemeriksaan menuju sasaran individual.')
label('Intervensi', 'Pantau elektrolit dan glukosa sesuai program. Amati kelemahan, kesemutan, kram, tetani, palpitasi, keringat dingin, atau perubahan kesadaran. Kolaborasikan koreksi dan insulin sesuai protokol; sinkronkan insulin dengan asupan dan laporkan bila enteral terhenti. Jika gangguan aktual terjadi, dokumentasikan kondisi aktual dan respons terapi.')
h('G Risiko infeksi')
label('Data dan luaran', 'Pertimbangkan alat invasif, kondisi umum, dan nekrosis bila ada. Targetnya tidak terjadi infeksi terkait pelayanan dan dugaan infeksi dikenali segera.')
label('Intervensi', 'Lakukan kebersihan tangan dan teknik aseptik; kaji infus, kateter, luka, serta drain. Tinjau kebutuhan alat setiap hari. Pantau suhu, kesadaran, perfusi, dan perburukan umum; kolaborasikan kultur serta terapi bila infeksi dicurigai.')
label('Rasional', 'Demam atau leukositosis dapat berasal dari inflamasi steril; antibiotik profilaksis tidak diberikan rutin hanya karena pankreatitis. [9]')
h('H Ansietas dan defisit pengetahuan')
label('Luaran dan intervensi', 'Pasien lebih tenang dan mampu menjelaskan kembali penyakit, terapi, serta tanda bahaya. Kaji kekhawatiran dan pemahaman, jelaskan tindakan dengan bahasa sederhana, libatkan keluarga dengan persetujuan pasien, gunakan teach-back, dan kolaborasikan dukungan psikologis atau ketergantungan alkohol bila diperlukan.')
h('I Kebutuhan tambahan')
p('Kaji intoleransi aktivitas, gangguan tidur, risiko jatuh, luka tekan, dan defisit perawatan diri. Terapkan bantuan aktivitas, mobilisasi bertahap, perawatan kulit, perubahan posisi individual, serta pencegahan tromboemboli sesuai penilaian medis dan protokol.')

page('8 Tanda bahaya dan kolaborasi')
p('Aktifkan bantuan medis atau tim respons cepat bila ditemukan perburukan. Jangan menunggu jadwal observasi berikutnya.')
table(['Temuan','Kekhawatiran klinis'],[
('Hipotensi, perfusi buruk, kesadaran menurun','Syok atau gangguan organ.'),
('Sesak, saturasi turun, oksigen meningkat','Gagal napas, komplikasi paru, atau kelebihan cairan.'),
('Urine turun nyata atau anuria','Gangguan perfusi, cedera ginjal, atau hambatan aliran urine.'),
('Abdomen tegang disertai oliguria atau sesak','Peningkatan tekanan intraabdomen.'),
('Nyeri mendadak hebat, perdarahan, Hb turun','Komplikasi intraabdomen atau perdarahan.'),
('Demam disertai perburukan hemodinamik','Infeksi atau sepsis.'),
('Kejang, aritmia, tetani, kesadaran berubah','Gangguan metabolik berat.')],[3.35,3.55])
p('Sambil menunggu bantuan, ulangi ABCDE, pastikan akses intravena, lakukan tindakan sesuai protokol, dan siapkan tren observasi serta cairan. Komunikasikan dengan SBAR: situasi, latar belakang, hasil penilaian, dan kebutuhan tindakan segera.')
p('Pada penyakit berat, tekanan intraabdomen dapat diukur menurut indikasi. Lingkar perut saja tidak menegakkan sindrom kompartemen abdomen. [7]')
h('Kolaborasi penyebab dan prosedur')
bullets([
'Pankreatitis bilier dengan kolangitis: koordinasikan evaluasi ERCP segera; ACG mendukung tindakan dalam 24 jam.',
'Pankreatitis bilier tanpa kolangitis: ERCP dini tidak dilakukan rutin.',
'Pankreatitis bilier ringan: koordinasikan kolesistektomi pada perawatan yang sama bila kondisi memungkinkan.',
'Nekrosis pankreas: keputusan tim multidisiplin; pada pasien stabil, tindakan biasanya ditunda sampai koleksi matang, sekitar empat minggu. Perburukan memerlukan evaluasi ulang segera.',
'Sebelum prosedur: verifikasi persiapan, persetujuan sesuai alur, alergi, obat, dan puasa yang diperlukan. Setelahnya pantau kesadaran, tanda vital, nyeri, serta perdarahan.'
])
p('Pemilihan dan waktu prosedur mengikuti urgensi serta stabilitas pasien. [4]')

page('9 Evaluasi dokumentasi dan persiapan pulang')
h('Evaluasi setiap sif dan setelah tindakan penting')
p('Bandingkan data aktual dengan nilai awal dan target. Nilai perfusi, keluaran urine, tanda kelebihan cairan, nyeri dan fungsi, oksigenasi, asupan serta toleransi, efek obat, infeksi, dan pemahaman pasien. Perbarui rencana bila target belum tercapai atau timbul masalah baru.')
label('S', 'Keluhan, skor nyeri, mual, toleransi makan, dan kekhawatiran pasien.')
label('O', 'Tanda vital, kesadaran, dukungan oksigen, urine, cairan, asupan, pemeriksaan fisik, dan hasil penunjang relevan.')
label('A', 'Diagnosis aktif dan status target: tercapai, sebagian tercapai, atau belum tercapai.')
label('P', 'Intervensi lanjutan, frekuensi pemantauan, kolaborasi, dan waktu evaluasi.')
p('Hindari catatan “keadaan membaik” tanpa data pendukung.')
h('Kesiapan pulang')
bullets([
'Kondisi klinis dan tanda vital stabil menurut penilaian bersama DPJP.',
'Nyeri terkendali dengan regimen yang dapat digunakan di rumah.',
'Makan/minum ditoleransi atau tersedia dukungan nutrisi yang aman.',
'Pasien dan keluarga mampu melakukan perawatan yang diperlukan.',
'Penyebab, pencegahan kekambuhan, resep, jadwal kontrol, kontak layanan, dan rencana tindakan lanjutan jelas.'
])
h('Edukasi pasien dan keluarga')
bullets([
'Ikuti rencana makan individual; hindari pembatasan berlebihan yang mengurangi kecukupan gizi.',
'Hindari alkohol, terutama bila menjadi penyebab, dan dapatkan dukungan berhenti bila diperlukan. Hentikan merokok.',
'Jelaskan dosis, durasi, efek samping, serta aturan penghentian obat sesuai resep.',
'Segera ke IGD bila nyeri memberat, muntah terus, tidak mampu makan/minum, demam, sesak, pingsan, kuning bertambah, atau urine sangat berkurang.',
'Laporkan penurunan berat badan, feses berminyak, sering haus, atau sering berkemih. Gunakan teach-back untuk memastikan pemahaman.'
])
p('Sertakan dukungan psikologis dan akses bantuan setelah pulang. [9] IAP 2025 menganjurkan skrining glukosa dan HbA1c mulai 3–6 bulan setelah episode, kemudian setiap 12 bulan; sesuaikan bila sudah terdapat diabetes atau gejala. [3]')

page('10 Daftar pustaka')
refs=[
('1','Cárdenas-Jaén K, dkk. Ibero-Latin American Guidelines on Acute Pancreatitis. United European Gastroenterology Journal. 2026.','https://doi.org/10.1002/ueg2.70190'),
('2','Pandanaboyana S, dkk. European Interdisciplinary Guidelines on Pain Management in Acute Pancreatitis. United European Gastroenterology Journal. 2026.','https://doi.org/10.1002/ueg2.70230'),
('3','IAP/APA/EPC/IPC/JPS Working Group. International Association of Pancreatology Revised Guidelines on Acute Pancreatitis 2025. Pancreatology. 2025;25:770–814.','https://doi.org/10.1016/j.pan.2025.04.020'),
('4','Tenner S, dkk. American College of Gastroenterology Guidelines: Management of Acute Pancreatitis. American Journal of Gastroenterology. 2024;119:419–437.','https://doi.org/10.14309/ajg.0000000000002645'),
('5','Arvanitakis M, dkk. ESPEN practical guideline on clinical nutrition in acute and chronic pancreatitis. Clinical Nutrition. 2024;43:395–412.','https://doi.org/10.1016/j.clnu.2023.12.019'),
('6','de-Madaria E, dkk. Aggressive or Moderate Fluid Resuscitation in Acute Pancreatitis. New England Journal of Medicine. 2022;387:989–1000.','https://doi.org/10.1056/NEJMoa2202884'),
('7','Leppäniemi A, dkk. 2019 WSES guidelines for the management of severe acute pancreatitis. World Journal of Emergency Surgery. 2019;14:27.','https://doi.org/10.1186/s13017-019-0247-0'),
('8','Banks PA, dkk. Classification of acute pancreatitis—2012: revision of the Atlanta classification and definitions by international consensus. Gut. 2013;62:102–111.','https://doi.org/10.1136/gutjnl-2012-302779'),
('9','NICE. Pancreatitis NG104. 2018; pembaruan 2020.','https://www.nice.org.uk/guidance/ng104/chapter/Recommendations')]
for n,t,u in refs:
    x=p(f'[{n}] {t}'); x.paragraph_format.keep_with_next=True
    x=p(u)
    for r in x.runs:r.font.size=Pt(9)
p('Rujukan terminologi: DPP PPNI. SDKI, SLKI, dan SIKI, edisi resmi yang berlaku di institusi, untuk validasi diagnosis, luaran, intervensi, dan dokumentasi keperawatan.')
d.core_properties.title='Panduan Asuhan Keperawatan Pasien Dewasa dengan Pankreatitis Akut'
d.core_properties.author=''
Path('output').mkdir(exist_ok=True)
path=Path('output/Panduan_Asuhan_Keperawatan_Pankreatitis_Akut.docx')
d.save(path)
print(path.resolve())
print('Jumlah kata:',sum(len(x.text.split()) for x in d.paragraphs))
